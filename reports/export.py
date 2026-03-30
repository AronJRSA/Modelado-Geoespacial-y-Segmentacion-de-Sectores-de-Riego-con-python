# Archivo: export.py
from docx import Document
from docx.shared import Inches
import os


def generar_reporte(
    output_path,
    imagen_mapa=None,
    imagen_3d=None,
    imagen_curvas=None,
    imagen_zonas=None
):

    doc = Document()

    # Título
    doc.add_heading(
        'Modelado Geoespacial y Segmentación de Sectores de Riego con Python',
        0
    )

    doc.add_paragraph('')
    doc.add_paragraph(
        'Procesamiento de datos KML y análisis topográfico para el modelado 3D '
        'del terreno, generación de curvas de nivel y segmentación de sectores '
        'de riego mediante técnicas de clustering.'
    )

    # ---- MAPA ----
    doc.add_paragraph('')
    doc.add_paragraph('Imagen satelital del predio', style='Heading 1')

    if imagen_mapa and os.path.exists(imagen_mapa):
        doc.add_picture(imagen_mapa, width=Inches(4.5))

    # ---- MODELO 3D ----
    doc.add_paragraph('')
    doc.add_paragraph('Modelo 3D', style='Heading 1')

    if imagen_3d and os.path.exists(imagen_3d):
        doc.add_picture(imagen_3d, width=Inches(4.5))

    doc.add_paragraph(
        'El modelo tridimensional representa la distribución de elevaciones del terreno.'
    )

    doc.add_paragraph('Zonas altas y bajas', style='List Bullet')
    doc.add_paragraph('Cambios en la pendiente', style='List Bullet')
    doc.add_paragraph('Estructura general del terreno', style='List Bullet')

    # ---- CURVAS ----
    doc.add_paragraph('')
    doc.add_paragraph('Curvas de nivel', style='Heading 1')

    if imagen_curvas and os.path.exists(imagen_curvas):
        doc.add_picture(imagen_curvas, width=Inches(4.5))

    doc.add_paragraph(
        'Las curvas representan líneas de igual altura y ayudan a analizar la topografía.'
    )

    doc.add_paragraph('Analizar la inclinación del terreno', style='List Bullet')
    doc.add_paragraph('Identificar zonas de escurrimiento', style='List Bullet')
    doc.add_paragraph('Planificación de riego', style='List Bullet')

    # ---- ZONAS ----
    doc.add_paragraph('')
    doc.add_paragraph('Segmentación de zonas', style='Heading 1')

    if imagen_zonas and os.path.exists(imagen_zonas):
        doc.add_picture(imagen_zonas, width=Inches(4.5))

    doc.add_paragraph(
        'La segmentación agrupa zonas con características similares para optimizar el riego.'
    )

    doc.add_paragraph('Sectores homogéneos', style='List Bullet')
    doc.add_paragraph('Gestión por zonas', style='List Bullet')
    doc.add_paragraph('Distribución eficiente del agua', style='List Bullet')

    # Guardar
    doc.save(output_path)