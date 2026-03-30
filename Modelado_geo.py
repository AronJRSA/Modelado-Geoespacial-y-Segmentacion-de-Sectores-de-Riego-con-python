# -*- coding: utf-8 -*-
"""
Created on Mon Mar 30 00:33:05 2026

@author: asron
"""


import os

proyecto = "Modelado"

carpetas = [
    "datos",
    "graficas",
    "reportes"
]

os.makedirs(proyecto, exist_ok=True)

for carpeta in carpetas:
    os.makedirs(f"{proyecto}/{carpeta}", exist_ok=True)
    
#librerias
import xml.etree.ElementTree as ET
import pandas as pd
import sys
import folium
import subprocess
import matplotlib.pyplot as plt
from sklearn.cluster import KMeans
import numpy as np
from scipy.interpolate import griddata
from docx import Document
from docx.shared import Inches

# CONFIGURACIÓN
archivo_kml = 'Modelado/datos/Huerta.kml'
archivo_salida = 'Modelado/datos/Levantamiento_Huerta_Aguacates.csv'
archivo_html = os.path.abspath('Modelado/graficas/mapa_temp.html')
archivo_png = os.path.abspath('Modelado/graficas/captura_huerta.png')

#PROCESAMIENTO KML
def procesar_kml():
    try:
        tree = ET.parse(archivo_kml)
        root = tree.getroot()
        ns = {'kml': 'http://www.opengis.net/kml/2.2'}
        datos_arboles = []
        for pm in root.findall('.//kml:Placemark', ns):
            nombre = pm.find('kml:name', ns).text if pm.find('kml:name', ns) is not None else "Sin_ID"
            coords = pm.find('.//kml:coordinates', ns)
            if coords is not None:
                c_list = coords.text.strip().split(',')
                datos_arboles.append({
                    'ID_Arbol': nombre,
                    'Longitud_X': float(c_list[0]),
                    'Latitud_Y': float(c_list[1]),
                    'Elevacion_Z': float(c_list[2])
                })
        df = pd.DataFrame(datos_arboles)
        df.to_csv(archivo_salida, index=False, encoding='utf-8')
        c_x, c_y = df['Longitud_X'].mean(), df['Latitud_Y'].mean()
        return df, c_y, c_x
    except Exception as e:
        print(f" Error KML: {e}")
        return None, None, None

#MAPA
def crear_mapa(df, centro_y, centro_x):
    mapa = folium.Map(location=[centro_y, centro_x], zoom_start=19,
                      tiles='https://mt1.google.com/vt/lyrs=s&x={x}&y={y}&z={z}', attr='Google')
    for _, fila in df.iterrows():
        folium.CircleMarker(location=[fila['Latitud_Y'], fila['Longitud_X']],
                            radius=2, color='lime', fill=True).add_to(mapa)
    mapa.save(archivo_html)

def realizar_captura_segura():
    """Crea un pequeño script de un solo uso para evitar el error de Spyder."""
    script_captura = "temp_capture_script.py"
    ruta_url = f"file:///{archivo_html.replace(os.sep, '/')}"
    
    codigo_script = f"""
import asyncio
from playwright.async_api import async_playwright

async def run():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={{'width': 1280, 'height': 720}})
        await page.goto(r'{ruta_url}')
        await asyncio.sleep(6)
        await page.screenshot(path=r'{archivo_png}')
        await browser.close()

if __name__ == '__main__':
    asyncio.run(run())
"""
    with open(script_captura, "w", encoding="utf-8") as f:
        f.write(codigo_script)
    

    subprocess.run([sys.executable, script_captura])
    

    if os.path.exists(script_captura):
        os.remove(script_captura) 

##modelado 3D
if __name__ == "__main__":
    df_puntos, c_y, c_x = procesar_kml()

    if df_puntos is not None:
        crear_mapa(df_puntos, c_y, c_x)
        realizar_captura_segura()
        fig = plt.figure(figsize=(12, 8))
        ax = fig.add_subplot(111, projection='3d')


        sc = ax.scatter(df_puntos['Longitud_X'], 
                        df_puntos['Latitud_Y'], 
                        df_puntos['Elevacion_Z'], 
                        c=df_puntos['Elevacion_Z'], 
                        cmap='terrain', s=60, edgecolors='k', alpha=0.8)

        ax.set_title('Modelo 3D de la Huerta', fontsize=15)
        ax.set_xlabel('Longitud (X)')
        ax.set_ylabel('Latitud (Y)')
        ax.set_zlabel('Elevación (msnm)')

        fig.colorbar(sc, ax=ax, label='Altura (msnm)', shrink=0.5, aspect=5)
        ax.view_init(elev=30, azim=45)
        plt.savefig("Modelado/graficas/3d.png",dpi=300) 
        plt.show()
        ## zonas de riego   
        X = df_puntos[['Longitud_X', 'Latitud_Y', 'Elevacion_Z']]
        ##Número de zonas (se puede cambiar a conveniencia)
        k = 4

        kmeans = KMeans(n_clusters=k, random_state=42)
        df_puntos['Zona'] = kmeans.fit_predict(X)
        #curvas de nivel
        x = df_puntos['Longitud_X'].values
        y = df_puntos['Latitud_Y'].values
        z = df_puntos['Elevacion_Z'].values
        xi = np.linspace(x.min(), x.max(), 100)
        yi = np.linspace(y.min(), y.max(), 100)

        xi, yi = np.meshgrid(xi, yi)
        zi = griddata((x, y), z, (xi, yi), method='cubic')
        plt.figure(figsize=(10,8))

        contour = plt.contour(xi, yi, zi, levels=15)
        plt.clabel(contour, inline=True, fontsize=8)
        plt.savefig("Modelado/graficas/curvas.png",dpi=300) 
        plt.scatter(x, y, c='red', s=10)

        plt.title('Curvas de nivel - Huerta')
        plt.xlabel('Longitud')
        plt.ylabel('Latitud')


        plt.figure(figsize=(10,8))


        plt.contourf(xi, yi, zi, levels=15, cmap='terrain')

        contour = plt.contour(xi, yi, zi, levels=15, colors='black', linewidths=0.5)


        plt.scatter(df_puntos['Longitud_X'], df_puntos['Latitud_Y'], c=df_puntos['Zona'], cmap='tab10', s=40)

        plt.title('Curvas + Zonas')
        plt.xlabel('Longitud')
        plt.ylabel('Latitud')

        plt.colorbar(label='Elevación')
        plt.savefig("Modelado/graficas/curvas_zona.png",dpi=300) 
        plt.show()

        #reporte
        doc = Document()
        doc.add_heading('Modelado Geoespacial y Segmentación de Sectores de Riego con Python', 0)
        doc.add_paragraph('')
        doc.add_paragraph('Procesamiento de datos KML y análisis topográfico para el modelado 3D del terreno, generación de curvas de nivel y segmentación de sectores de riego mediante técnicas de clustering.')
        doc.add_paragraph('')
        doc.add_paragraph('Imagen satelital del predio', style='Heading 1')
        doc.add_paragraph('')
        if os.path.exists('Modelado/graficas/captura_huerta.png'):
            doc.add_picture('Modelado/graficas/captura_huerta.png', width=Inches(4.5))
        doc.add_paragraph('')
        doc.add_paragraph('Modelo 3D', style='Heading 1')
        doc.add_paragraph('')
        if os.path.exists('Modelado/graficas/3d.png'):
            doc.add_picture('Modelado/graficas/3d.png', width=Inches(4.5))
        doc.add_paragraph('')
        doc.add_paragraph('El modelo tridimensional representa la distribución de elevaciones del terreno, permitiendo una visualización más intuitiva de la topografía.')
        doc.add_paragraph('Facilita la identificación de:')
        doc.add_paragraph('')
        doc.add_paragraph('Zonas altas y bajas', style='List Bullet')
        doc.add_paragraph('Cambios en la pendiente', style='List Bullet')
        doc.add_paragraph('Estructura general del terreno',style= 'List Bullet')
        doc.add_paragraph('Curvas de nivel', style='Heading 1')
        doc.add_paragraph('')
        if os.path.exists('Modelado/graficas/curvas.png'):
            doc.add_picture('Modelado/graficas/curvas.png', width=Inches(4.5))
        doc.add_paragraph('')
        doc.add_paragraph('En la imagen se muestran curvas de nivel generadas a partir de la interpolación de elevaciones del terreno')
        doc.add_paragraph('Estas curvas representan líneas de igual altura y permiten identificar la pendiente y variación topográfica del predio.')
        doc.add_paragraph('Son útiles para:')
        doc.add_paragraph('')
        doc.add_paragraph('Analizar la inclinación del terreno', style='List Bullet')
        doc.add_paragraph('Identificar zonas de escurrimiento', style='List Bullet')
        doc.add_paragraph('Apoyar la planificación de riego y drenaje', style='List Bullet')
        doc.add_paragraph('')
        doc.add_paragraph('Segmentación de zonas', style='Heading 1')
        doc.add_paragraph('')
        if os.path.exists('Modelado/graficas/curvas_zona.png'):
            doc.add_picture('Modelado/graficas/curvas_zona.png', width=Inches(4.5))
        doc.add_paragraph('')
        doc.add_paragraph('La imagen muestra la segmentación del terreno en diferentes zonas mediante técnicas de clustering')
        doc.add_paragraph('Cada zona agrupa puntos con características espaciales y topográficas similares.')
        doc.add_paragraph('')
        doc.add_paragraph('Este tipo de segmentación permite:')
        doc.add_paragraph('')
        doc.add_paragraph('Dividir el terreno en sectores homogéneos', style='List Bullet')
        doc.add_paragraph('Facilitar la gestión del riego por zonas', style='List Bullet')
        doc.add_paragraph('Reducir variaciones en la distribución del agua', style='List Bullet')
        doc.save("Modelado/reportes/Reporte_Final_Consultoria_24.docx")
        print('listo')